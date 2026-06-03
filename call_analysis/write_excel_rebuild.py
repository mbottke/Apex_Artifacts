"""Write the V2 full-rebuild workbook: rewrite every calendar cell from the solved schedule,
preserving all cell formatting and the live tracking formulas (which recompute on open)."""
import openpyxl, json, datetime, sys, os
WHICH=sys.argv[1] if len(sys.argv)>1 else 'rebuilt'
OUT=sys.argv[2] if len(sys.argv)>2 else 'deliverables/26-27_On_Call_Schedule_v2_full_rebuild.xlsx'
SRC="/root/.claude/uploads/8fdbdbc1-c81a-469d-bd88-b325324f7770/10a2b575-202627_On_Call_Schedule.xlsx"
B='/home/user/Apex_Artifacts/call_analysis/'
S=json.load(open(B+f'{WHICH}.json'))['sched']
months=["July","August","September","October","November","December",
        "January","February","March","April","May","June"]
DAY_COLS=[3,5,7,9,11,15,23]
import datetime as _dt
wb_vals=openpyxl.load_workbook(SRC,data_only=True)
wb=openpyxl.load_workbook(SRC,data_only=False)
# date -> authoritative (tab,row,col) using cached dates
cellmap={}
for m in months:
    ws=wb_vals[m]
    for col in DAY_COLS:
        cur=None
        for row in range(4,24):
            v=ws.cell(row=row,column=col).value
            if isinstance(v,(_dt.datetime,_dt.date)) and not isinstance(v,bool):
                cur=v.date() if isinstance(v,_dt.datetime) else v
            elif isinstance(v,str) and v.strip() and cur is not None:
                ds=cur.isoformat(); inm=cur.strftime('%B')==m
                if ds not in cellmap or (inm and not cellmap[ds][3]): cellmap[ds]=(m,row,col,inm)
written=0; miss=[]
for ds,e in S.items():
    if ds not in cellmap: miss.append(ds); continue
    m,row,col,_=cellmap[ds]
    line1='/'.join(e['interns']); line2='/'.join(e['seniors'])
    val=(line1+'\n'+line2) if line1 else line2
    wb[m].cell(row=row,column=col).value=val
    written+=1
wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
os.makedirs(os.path.dirname('/home/user/Apex_Artifacts/'+OUT),exist_ok=True)
wb.save('/home/user/Apex_Artifacts/'+OUT)
print(f"[{WHICH}] cells written={written} missing={len(miss)} -> {OUT}")
if miss: print("  missing dates:",miss[:20])
