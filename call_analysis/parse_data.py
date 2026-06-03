"""Parse the Rotations Master Schedule and the On-Call Schedule into structured JSON."""
import openpyxl, json, datetime, re

ROT = "/root/.claude/uploads/8fdbdbc1-c81a-469d-bd88-b325324f7770/a67f435c-Rotations_20262027.xlsx"
SCHED = "/root/.claude/uploads/8fdbdbc1-c81a-469d-bd88-b325324f7770/10a2b575-202627_On_Call_Schedule.xlsx"

PGY = {
    'Hard':'R1','Oehm':'R1','Schutt':'R1','Stanek':'R1','Strand':'R1','Wohlgemuth':'R1',
    'Beutel':'R2','Bottke':'R2','Bower':'R2','Gaspar':'R2','Kendrick':'R2','Sublette':'R2','Patel':'R2',
    'Johnson':'R3','Lux':'R3','Mautino':'R3','Mudondo':'R3',
}
RESIDENT_ROWS = {  # Master Schedule top/bottom rows
    'Johnson':(46,47),'Lux':(48,49),'Mautino':(50,51),'Mudondo':(52,53),
    'Beutel':(58,59),'Bottke':(60,61),'Bower':(62,63),'Gaspar':(64,65),
    'Kendrick':(66,67),'Sublette':(68,69),'Patel':(70,71),
    'Wohlgemuth':(76,77),'Schutt':(78,79),'Hard':(80,81),'Oehm':(82,83),
    'Strand':(84,85),'Stanek':(86,87),
}

def parse_block(s):
    # "7/5/26-7/18/26" -> (date,date)
    a,b = s.strip().split('-')
    def pd(x):
        m,d,y = x.strip().split('/')
        return datetime.date(2000+int(y), int(m), int(d))
    return pd(a), pd(b)

wb = openpyxl.load_workbook(ROT, data_only=True)
ws = wb["Master Schedule"]

# Block date ranges from rows 44 (top) and 45 (bottom)
DATA_COLS = list(range(2,9)) + list(range(10,17))  # B..H, J..P
top_blocks, bot_blocks = {}, {}
for c in DATA_COLS:
    t = ws.cell(row=44, column=c).value
    b = ws.cell(row=45, column=c).value
    if t: top_blocks[c] = parse_block(str(t))
    if b: bot_blocks[c] = parse_block(str(b))

# Parse rotations: resident -> list of {start,end,rotation}
rotations = {}
for name,(tr,br) in RESIDENT_ROWS.items():
    spans = []
    for c in DATA_COLS:
        tv = ws.cell(row=tr, column=c).value
        bv = ws.cell(row=br, column=c).value
        tv = str(tv).strip() if tv is not None and str(tv).strip() else None
        bv = str(bv).strip() if bv is not None and str(bv).strip() else None
        if tv is None and bv is None:
            continue
        if tv and bv:
            s,e = top_blocks[c]; spans.append((s,e,tv))
            s2,e2 = bot_blocks[c]; spans.append((s2,e2,bv))
        elif tv and not bv:
            # 4-week rotation spanning full column (or single block for H,P)
            s = top_blocks[c][0]
            e = bot_blocks[c][1] if c in bot_blocks else top_blocks[c][1]
            spans.append((s,e,tv))
        elif bv and not tv:
            s,e = bot_blocks[c]; spans.append((s,e,bv))
    rotations[name] = [(s.isoformat(), e.isoformat(), r) for s,e,r in spans]

# Build daily rotation lookup
def rotation_on(name, d):
    for s,e,r in rotations[name]:
        sd = datetime.date.fromisoformat(s); ed = datetime.date.fromisoformat(e)
        if sd <= d <= ed:
            return r
    return None

# ---- Parse call assignments from all 12 months ----
wbs = openpyxl.load_workbook(SCHED, data_only=True)
months = ["July","August","September","October","November","December",
          "January","February","March","April","May","June"]
DAY_COLS = [3,5,7,9,11,15,23]  # C,E,G,I,K,O,W = Sun..Sat
ALLNAMES = set(PGY.keys())

def split_names(line):
    # split "Wohlgemuth/Patel" or "Schutt/Patel" -> list; strip spaces
    parts = re.split(r'[\/,&]', line)
    out = []
    for p in parts:
        p = p.strip()
        if not p: continue
        # match to known resident by surname token
        matched = None
        for nm in ALLNAMES:
            if nm.lower() == p.lower() or p.lower().startswith(nm.lower()) or nm.lower() in p.lower():
                matched = nm; break
        out.append(matched if matched else p)
    return out

assignments = []  # {date, dow, interns:[], seniors:[], raw}
for m in months:
    ws = wbs[m]
    for col in DAY_COLS:
        current_date = None
        for row in range(4, 24):
            cell = ws.cell(row=row, column=col)
            v = cell.value
            if v is None: continue
            if isinstance(v, datetime.datetime):
                current_date = v.date()
            elif isinstance(v, datetime.date):
                current_date = v
            elif isinstance(v, str) and v.strip():
                # assignment cell
                if current_date is None: continue
                lines = [ln.strip() for ln in v.split('\n') if ln.strip()]
                interns, seniors = [], []
                if len(lines) >= 1:
                    for nm in split_names(lines[0]):
                        (interns if PGY.get(nm)=='R1' else seniors).append(nm)
                if len(lines) >= 2:
                    for nm in split_names(lines[1]):
                        (interns if PGY.get(nm)=='R1' else seniors).append(nm)
                assignments.append({
                    'date': current_date.isoformat(),
                    'dow': current_date.strftime('%a'),
                    'month_tab': m,
                    'line1': lines[0] if lines else '',
                    'line2': lines[1] if len(lines)>1 else '',
                    'interns': interns,
                    'seniors': seniors,
                    'raw': v.replace('\n','/'),
                })

# Deduplicate (months overlap at boundaries: e.g., July tab shows some Aug days)
seen = {}
for a in assignments:
    # prefer the tab whose month matches the date's month
    d = datetime.date.fromisoformat(a['date'])
    a['in_tab_month'] = (d.strftime('%B') == a['month_tab'])
    key = a['date']
    if key not in seen or (a['in_tab_month'] and not seen[key]['in_tab_month']):
        seen[key] = a
assign_dedup = sorted(seen.values(), key=lambda x: x['date'])

out = {
    'rotations': rotations,
    'pgy': PGY,
    'assignments': assign_dedup,
    'blocks_top': {c:[top_blocks[c][0].isoformat(),top_blocks[c][1].isoformat()] for c in top_blocks},
    'blocks_bot': {c:[bot_blocks[c][0].isoformat(),bot_blocks[c][1].isoformat()] for c in bot_blocks},
}
with open('/home/user/Apex_Artifacts/call_analysis/data.json','w') as f:
    json.dump(out, f, indent=1, default=str)

print("Residents:", len(rotations))
print("Call days parsed (deduped):", len(assign_dedup))
print("Date range:", assign_dedup[0]['date'], "to", assign_dedup[-1]['date'])
# sanity: print rotation spans for a couple residents
for nm in ['Gaspar','Stanek','Bottke','Johnson']:
    print(f"\n{nm} ({PGY[nm]}) rotations:")
    for s,e,r in rotations[nm]:
        print(f"   {s} -> {e}: {r}")
