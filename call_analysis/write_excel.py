"""Write a revised On-Call Schedule .xlsx in the SAME tab format as the original.
Only the changed calendar cells are edited (token-level surname replacement); all tracking
tables and the YTD tab are live formulas and recompute automatically on open."""
import openpyxl, json, datetime, sys, os

WHICH = sys.argv[1] if len(sys.argv)>1 else 'repaired'
OUT   = sys.argv[2] if len(sys.argv)>2 else 'deliverables/26-27_On_Call_Schedule_v1_minimal_repair.xlsx'
SRC = "/root/.claude/uploads/8fdbdbc1-c81a-469d-bd88-b325324f7770/10a2b575-202627_On_Call_Schedule.xlsx"
BASE='/home/user/Apex_Artifacts/call_analysis/'
D=json.load(open(BASE+'data.json')); A0=D['assignments']; PGY=D['pgy']
R=json.load(open(BASE+f'{WHICH}.json')); sched=R['sched']
orig={a['date']:a for a in A0}

months=["July","August","September","October","November","December",
        "January","February","March","April","May","June"]
DAY_COLS=[3,5,7,9,11,15,23]

import datetime as _dt
# Calendar date cells are FORMULAS (=C5, =G5+1) -> resolve dates from a cached-value load,
# then EDIT the formula-preserving load at the same coordinates.
wb_vals=openpyxl.load_workbook(SRC, data_only=True)
wb=openpyxl.load_workbook(SRC, data_only=False)  # preserve formulas + formatting (this one is saved)

# build date -> authoritative (tab,row,col) for assignment cells holding NAMES
cellmap={}
for m in months:
    ws=wb_vals[m]
    for col in DAY_COLS:
        cur=None
        for row in range(4,24):
            v=ws.cell(row=row,column=col).value
            if isinstance(v,(_dt.datetime,_dt.date)) and not isinstance(v,bool):
                cur=v.date() if isinstance(v,_dt.datetime) else v
            elif isinstance(v,str) and v.strip() and cur is not None:
                ds=cur.isoformat()
                inmonth = cur.strftime('%B')==m
                if ds not in cellmap or (inmonth and not cellmap[ds][3]):
                    cellmap[ds]=(m,row,col,inmonth)

def replace_token(text, old_sn, new_sn):
    """Replace the token containing surname old_sn with new_sn, preserving line/slash structure & dropping markers."""
    out_lines=[]
    for line in text.split('\n'):
        toks=line.split('/')
        for i,t in enumerate(toks):
            if old_sn.lower() in t.strip().lower():
                # preserve leading/trailing spaces minimally; drop markers like ** by replacing whole token
                lead=' '*(len(t)-len(t.lstrip()));
                toks[i]=lead+new_sn
        out_lines.append('/'.join(toks))
    return '\n'.join(out_lines)

edits=[]
for ds, e in sched.items():
    o=orig.get(ds)
    if not o: continue
    for role in ('interns','seniors'):
        old=list(o[role]); new=list(e[role])
        removed=[x for x in old if x not in new]
        added=[x for x in new if x not in old]
        for rem,add in zip(removed,added):
            edits.append((ds,rem,add))

# apply edits
applied=0; problems=[]
for ds,rem,add in edits:
    if ds not in cellmap: problems.append((ds,rem,add,'no cell')); continue
    m,row,col,_=cellmap[ds]
    ws=wb[m]; cell=ws.cell(row=row,column=col)
    txt=cell.value
    if not isinstance(txt,str) or rem.lower() not in txt.lower():
        problems.append((ds,rem,add,f'token not found in {m}!{cell.coordinate}={txt!r}')); continue
    cell.value=replace_token(txt,rem,add)
    applied+=1

wb.calculation.fullCalcOnLoad = True   # force Excel/LibreOffice to recompute all formulas on open
wb.calculation.forceFullCalc = True
os.makedirs(os.path.dirname('/home/user/Apex_Artifacts/'+OUT),exist_ok=True)
wb.save('/home/user/Apex_Artifacts/'+OUT)
print(f"[{WHICH}] edits={len(edits)} applied={applied} problems={len(problems)}")
for p in problems: print("   PROBLEM:",p)
print("Saved ->",OUT)
